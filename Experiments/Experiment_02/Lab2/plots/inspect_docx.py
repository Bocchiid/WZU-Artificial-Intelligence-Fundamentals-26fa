# -*- coding: utf-8 -*-
"""检查 docx 模板结构：表格布局 + 正文字体字号，输出为可读文本。"""

import sys
from docx import Document
from docx.shared import Pt


def run_fmt(r):
    f = r.font
    nm = f.name
    sz = f.size.pt if f.size else None
    east = None
    try:
        rpr = r._element.rPr
        if rpr is not None and rpr.rFonts is not None:
            east = rpr.rFonts.get(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
    except Exception:
        pass
    return "%s/%s/%s" % (nm or "-", east or "-", ("%.0fpt" % sz) if sz else "-")


def main(path, out):
    d = Document(path)

    L = []
    L.append("=" * 78)
    L.append("文档: %s" % path)
    L.append("=" * 78)

    L.append("\n【正文段落】共 %d 段" % len(d.paragraphs))
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if t:
            L.append("  p%-3d [%s] %s" % (i, p.style.name, t[:60]))
            for r in p.runs[:3]:
                L.append("         run: %r  <%s>" % (r.text[:30], run_fmt(r)))

    L.append("\n【默认样式】")
    try:
        st = d.styles['Normal']
        L.append("  Normal font: %s  size: %s" % (st.font.name, st.font.size))
    except Exception as e:
        L.append("  (读取失败 %s)" % e)

    L.append("\n【表格】共 %d 个" % len(d.tables))
    for ti, tb in enumerate(d.tables):
        L.append("  --- 表 %d: %d 行 ---" % (ti + 1, len(tb.rows)))
        for ri, row in enumerate(tb.rows):
            cells = row.cells
            # 去重相邻重复（合并单元格会重复出现）
            seen, uniq = None, []
            for c in cells:
                if c._tc is not seen:
                    uniq.append(c)
                    seen = c._tc
            txts = [c.text.strip().replace("\n", "\\n")[:34] for c in uniq]
            L.append("    行%-3d %d格: %s" % (ri + 1, len(uniq), " | ".join(txts)))

    # 表格里正文的字体
    L.append("\n【表格内文字格式抽样】")
    for ti, tb in enumerate(d.tables):
        for ri, row in enumerate(tb.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            L.append("  表%d 行%d [%s] %r <%s>"
                                     % (ti + 1, ri + 1, p.style.name,
                                        r.text.strip()[:26], run_fmt(r)))
                            break
                    else:
                        continue
                    break
                else:
                    continue
                break
            else:
                continue
            break

    txt = "\n".join(L)
    with open(out, "w", encoding="utf-8") as f:
        f.write(txt)
    print("written:", out, len(txt), "chars")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
